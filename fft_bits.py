"""ビット列を一定のビット数に分割し、FFTで解析してExcelに保存する。"""

# 実行コマンド（WSLのターミナルで実行）:
# cd /path/to/fft-bit-analysis
# 初回・ライブラリ追加時:
# ./.venv-wsl/bin/python -m pip install -r requirements.txt
# bin_files内のbinファイルが1つだけの場合:
# ./.venv-wsl/bin/python fft_bits.py
# 入力ファイルを指定する場合（bin_files内のファイル名）:
# ./.venv-wsl/bin/python fft_bits.py bits.bin
# 出力先も指定する場合:
# ./.venv-wsl/bin/python fft_bits.py bits.bin --output fft_results
# 上記の入力に対応するExcel出力: fft_results/bits.xlsx

import argparse
import csv
import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

INPUT_DIR = Path(__file__).resolve().parent / "bin_files"


def write_excel(path, summary, rows, frequencies, average_psd):
    """大きなスペクトルも行単位で書き出し、入力元をExcel内に記録する。"""
    workbook = Workbook(write_only=True)

    def sheet(name, headers, widths):
        target = workbook.create_sheet(name)
        target.freeze_panes = "A2"
        for column, width in widths.items():
            target.column_dimensions[column].width = width
        cells = []
        for label in headers:
            cell = WriteOnlyCell(target, value=label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24476A")
            cells.append(cell)
        target.append(cells)
        return target

    overview = sheet("Summary", ["項目", "値"], {"A": 44, "B": 100})
    for key, value in summary.items():
        if isinstance(value, dict):
            for metric, number in value.items():
                overview.append([f"{key}.{metric}", number])
        else:
            # パスは先頭が '=' の場合も数式として扱わない。
            if isinstance(value, str):
                cell = WriteOnlyCell(overview, value=value)
                cell.data_type = "s"
                value = cell
            overview.append([key, value])

    def write_table(name, headers, values, widths):
        target = None
        # Excelの1シートの上限はヘッダーを含めて1,048,576行。
        for index, row in enumerate(values):
            if index % 1_048_575 == 0:
                part = index // 1_048_575 + 1
                target = sheet(name if part == 1 else f"{name}_{part}", headers, widths)
            target.append(row)

    write_table("Peaks", ["chunk", "start_seconds", "fraction_ones",
                          "peak_frequency_hz", "peak_psd_bit_squared_per_hz"],
                rows, {"A": 12, "B": 22, "C": 22, "D": 26, "E": 36})
    write_table("Average spectrum", ["frequency_hz", "mean_psd_bit_squared_per_hz"],
                ((float(frequency), float(psd)) for frequency, psd in zip(frequencies, average_psd)),
                {"A": 24, "B": 36})
    workbook.save(path)


def resolve_input(input_path=None):
    """省略時は専用フォルダの唯一のbinファイルを選択する。"""
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    if input_path is not None:
        if not input_path.is_absolute() and input_path.parent == Path("."):
            input_path = INPUT_DIR / input_path
        if not input_path.is_file():
            raise ValueError(f"入力ファイルがありません: {input_path}")
        return input_path
    files = sorted(p for p in INPUT_DIR.iterdir() if p.is_file() and p.suffix.lower() == ".bin")
    if not files:
        raise ValueError(f"{INPUT_DIR} に解析するbinファイルを置いてください。")
    if len(files) > 1:
        raise ValueError("binファイルが複数あります。ファイル名を指定してください: "
                         + ", ".join(p.name for p in files))
    return files[0]


def bit_chunks(path, size, input_format="bytes", bit_order="big", max_bits=None):
    """保存形式に応じて0/1サンプルへ変換し、指定ビット数ずつ読み込む。"""
    if input_format not in ("bytes", "packed", "text"):
        raise ValueError(f"未対応の入力形式: {input_format}")
    pending = bytearray()
    remaining = max_bits
    with path.open("rb") as source:
        while remaining is None or remaining > 0:
            read_size = 1024 * 1024
            if remaining is not None:
                read_size = min(read_size, (remaining + 7) // 8 if input_format == "packed" else remaining)
            block = source.read(read_size)
            if not block:
                break
            if input_format == "text":
                block = block.translate(None, b" \t\r\n\v\f")
                if block.translate(None, b"01"):
                    raise ValueError("入力には0/1と空白・改行だけを使用してください（ASCII/UTF-8、BOMなし）。")
                block = (np.frombuffer(block, dtype=np.uint8) - 48).tobytes()
            elif input_format == "packed":
                block = np.unpackbits(np.frombuffer(block, dtype=np.uint8), bitorder=bit_order).tobytes()
            elif block.translate(None, b"\x00\x01"):
                raise ValueError("bin入力は1バイトごとに数値0または1を保存してください。別形式の場合は--input-formatを指定してください。")
            if remaining is not None:
                block = block[:remaining]
                remaining -= len(block)
            pending.extend(block)
            offset = 0
            while len(pending) - offset >= size:
                yield np.frombuffer(bytes(pending[offset:offset + size]), dtype=np.uint8).astype(np.float64)
                offset += size
            del pending[:offset]
    if pending:
        raise ValueError(f"末尾に{len(pending):,} bitの端数があります。区間サイズを確認してください。")


def analyze(input_path, output_dir, chunk_bits=1_000_000, chunks=1000,
            sample_rate=100_000_000.0, window_name="hann", input_format="bytes", bit_order="big"):
    if chunk_bits < 3 or chunks < 1 or not np.isfinite(sample_rate) or sample_rate <= 0:
        raise ValueError("区間サイズは3以上、区間数は1以上、サンプリング周波数は正の有限値にしてください。")
    window = np.hanning(chunk_bits) if window_name == "hann" else np.ones(chunk_bits)
    scale = sample_rate * np.dot(window, window)
    frequencies = np.fft.rfftfreq(chunk_bits, d=1 / sample_rate)
    psd_sum = np.zeros(len(frequencies))
    rows = []
    count = 0
    # 指定サンプル数に達したら終了し、それ以降のパディングを除外する。
    for count, values in enumerate(bit_chunks(input_path, chunk_bits, input_format, bit_order,
                                             max_bits=chunk_bits * chunks), 1):
        # 直流成分を除去し、区間端の影響を抑える窓を掛けてFFTする。
        fraction_ones = float(values.mean())
        values -= fraction_ones
        transform = np.fft.rfft(values * window)
        psd = np.abs(transform) ** 2 / scale
        # 片側PSDへ変換する。直流と偶数長のナイキスト成分は2倍しない。
        psd[1:-1 if chunk_bits % 2 == 0 else None] *= 2
        psd_sum += psd
        # 0 Hzを除いて最大ピークを探す。全0・全1の区間は欠損扱いにする。
        valid = bool(np.any(psd[1:] > 0))
        index = int(np.argmax(psd[1:])) + 1 if valid else None
        rows.append((count,
                    (count - 1) * chunk_bits / sample_rate,
                    fraction_ones,
                    float(frequencies[index]) if valid else None,
                    float(psd[index]) if valid else None))
        if count == 1 or count % 25 == 0 or count == chunks:
            print(f"処理済み: {count}/{chunks}区間", flush=True)
    if count != chunks:
        raise ValueError(f"入力は{count}区間です。指定の{chunks}区間に足りません。")

    # 各区間のパワースペクトルを平均する（複素FFTの平均ではない）。
    average_psd = psd_sum / count
    peaks = np.array([row[3] for row in rows if row[3] is not None])
    powers = np.array([row[4] for row in rows if row[4] is not None])
    average_index = int(np.argmax(average_psd[1:])) + 1 if np.any(average_psd[1:] > 0) else None

    def statistics(data):
        # 有効な区間のピークについて、平均・母分散・標準偏差を求める。
        return {"mean": float(data.mean()) if data.size else None,
                "variance": float(data.var(ddof=0)) if data.size else None,
                "std": float(data.std(ddof=0)) if data.size else None}

    summary = {
        "input": str(input_path.resolve()),
        "sample_rate_hz": sample_rate,
        "chunk_bits": chunk_bits,
        "chunks": count,
        "total_bits": chunk_bits * count,
        "window": window_name,
        "frequency_bin_hz": sample_rate / chunk_bits,
        "valid_peak_chunks": int(peaks.size),
        "peak_frequency_hz": statistics(peaks),
        "peak_psd_bit_squared_per_hz": statistics(powers),
        "average_spectrum_peak_frequency_hz": float(frequencies[average_index]) if average_index is not None else None,
        "variance_ddof": 0,
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix = input_path.stem
    excel_path = output_dir / f"{prefix}.xlsx"
    write_excel(excel_path, summary, rows, frequencies, average_psd)
    with (output_dir / f"{prefix}_peaks.csv").open("w", newline="", encoding="utf-8") as target:
        writer = csv.writer(target)
        writer.writerow(["chunk", "start_seconds", "fraction_ones", "peak_frequency_hz", "peak_psd_bit_squared_per_hz"])
        writer.writerows(rows)
    np.savetxt(output_dir / f"{prefix}_average_spectrum.csv",
                np.column_stack((frequencies, average_psd)), delimiter=",",
                header="frequency_hz,mean_psd_bit_squared_per_hz", comments="")
    (output_dir / f"{prefix}_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(f"Excel出力: {excel_path.resolve()}", flush=True)
    return summary


def main():
    parser = argparse.ArgumentParser(description="ビット列を分割してFFTし、ピーク周波数の平均・分散と平均パワースペクトルを保存します。")
    parser.add_argument("input", type=Path, nargs="?", help="bin_files内のファイル名、またはファイルパス。省略時は唯一のbinファイルを使用")
    parser.add_argument("--input-format", choices=["bytes", "packed", "text"], default="bytes")
    parser.add_argument("--bit-order", choices=["big", "little"], default="big")
    parser.add_argument("--output", type=Path, default=Path("fft_results"))
    parser.add_argument("--chunk-bits", type=int, default=1_000_000)
    parser.add_argument("--chunks", type=int, default=1000,
                        help="解析区間数。既定は1 Mbit × 1000 = 1 Gbit。以降のデータは無視")
    parser.add_argument("--sample-rate", type=float, default=100_000_000.0, help="Hz")
    parser.add_argument("--window", choices=["hann", "none"], default="hann")
    args = parser.parse_args()
    try:
        input_path = resolve_input(args.input)
        print(f"入力ファイル: {input_path.resolve()}", flush=True)
        summary = analyze(input_path, args.output, args.chunk_bits, args.chunks,
                          args.sample_rate, args.window, args.input_format, args.bit_order)
    except (OSError, ValueError) as error:
        parser.exit(1, f"エラー: {error}\n")
    print(json.dumps(summary, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
