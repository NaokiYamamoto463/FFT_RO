"""入力名とExcel出力・CSVの対応を小さな既知信号で検証する。"""
import csv
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from fft_bits import analyze


class ExcelOutputTests(unittest.TestCase):
    def test_input_names_and_values(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "results"
            first_content = None
            for name in ("bits_01.bin", "測定 02.v1.bin"):
                source = root / name
                source.write_bytes(bytes([0, 0, 1, 1]) * 4 + bytes(8))
                summary = analyze(source, output, chunk_bits=8, chunks=3,
                                  sample_rate=800, window_name="none")
                prefix = source.stem
                book = load_workbook(output / f"{prefix}.xlsx", read_only=True)
                try:
                    self.assertEqual(book.sheetnames, ["Summary", "Peaks", "Average spectrum"])
                    overview = dict(list(book["Summary"].values)[1:])
                    self.assertEqual(overview["input"], str(source.resolve()))
                    self.assertEqual(overview["valid_peak_chunks"], 2)
                    self.assertEqual(overview["peak_frequency_hz.mean"], 200)
                    peaks = list(book["Peaks"].values)
                    self.assertEqual(len(peaks), 4)
                    self.assertEqual(peaks[1][3], 200)
                    self.assertEqual(peaks[3][:3], (3, 0.02, 0))
                    self.assertTrue(all(value is None for value in peaks[3][3:]))
                    with (output / f"{prefix}_average_spectrum.csv").open() as stream:
                        expected = list(csv.reader(stream))
                    actual = list(book["Average spectrum"].values)
                    self.assertEqual(actual[0], tuple(expected[0]))
                    self.assertEqual(len(actual), len(expected))
                    for excel_row, csv_row in zip(actual[1:], expected[1:]):
                        for value, text in zip(excel_row, csv_row):
                            self.assertAlmostEqual(value, float(text), places=14)
                    self.assertEqual(json.loads((output / f"{prefix}_summary.json").read_text()), summary)
                    self.assertEqual(book["Peaks"].cell(2, 4).data_type, "n")
                finally:
                    book.close()
                first_path = output / "bits_01.xlsx"
                if first_content is None:
                    first_content = first_path.read_bytes()
                else:
                    self.assertEqual(first_path.read_bytes(), first_content)
            self.assertEqual(len(list(output.iterdir())), 8)


if __name__ == "__main__":
    unittest.main()
